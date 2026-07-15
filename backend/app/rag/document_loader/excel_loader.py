"""Excel 文档解析器"""
import logging
from openpyxl import load_workbook
from app.rag.document_loader.base import BaseDocumentLoader
from app.rag.models import LoadedDocument, DocumentType

logger = logging.getLogger(__name__)

class ExcelLoader(BaseDocumentLoader):
    def __init__(self):
        super().__init__()
        self.supported_types = [DocumentType.EXCEL]
    
    def can_handle(self, file_path: str, mime_type: str = "") -> bool:
        ext = file_path.lower().split('.')[-1]
        return ext in ('xls', 'xlsx', 'csv') or 'excel' in mime_type.lower() or 'spreadsheet' in mime_type.lower()
    
    async def load(self, file_path: str, **kwargs) -> LoadedDocument:
        wb = load_workbook(file_path, data_only=True)
        tables = []
        content_parts = []
        for sheet in wb.sheetnames:
            ws = wb[sheet]
            rows = []
            for row in ws.iter_rows(values_only=True):
                if any(cell is not None for cell in row):
                    rows.append([str(cell) if cell is not None else "" for cell in row])
            if rows:
                tables.append({"sheet": sheet, "data": rows})
                # 也将表格转为文本
                header = rows[0] if rows else []
                for row in rows[1:]:
                    row_text = " | ".join(f"{h}: {v}" for h, v in zip(header, row) if v)
                    if row_text:
                        content_parts.append(f"[{sheet}] {row_text}")
        content = "\n".join(content_parts) if content_parts else ""
        return LoadedDocument(
            source_type=DocumentType.EXCEL,
            file_path=file_path,
            file_name=file_path.split('/')[-1].split('\\')[-1],
            content=content,
            tables=tables,
            metadata={"sheet_count": len(wb.sheetnames), "sheet_names": list(wb.sheetnames)},
        )
