"""项目级「需求分析结果」：结构化、可追溯、不做测试设计。"""
from __future__ import annotations

import asyncio
import concurrent.futures
import io
import json
import re
from datetime import datetime, timezone
from typing import Any, Optional

from app.llm import get_gateway
from app.services.project_memory import ProjectMemoryService

DOC_TITLE = "需求分析"
DOC_KIND = "understanding"
DOC_WORKSPACE = "understand"
MAX_SOURCE_CHARS = 80_000
EVIDENCE_KINDS = {"explicit", "inferred", "missing"}
RISK_LEVELS = {"高", "中", "低"}

SYSTEM_PROMPT = """你是专业测试需求分析师，不是聊天机器人，也不是测试用例设计师。

任务：把测试人员提供的需求材料，分析成一份结构化、可测试、可追踪的「需求分析结果」。
只回答「这个系统/功能到底要实现什么」，不要回答「应该怎么测」，禁止输出完整测试用例、测试脚本或执行任务。

硬性规则：
1. 只根据给定原文分析。禁止编造不存在的业务规则、字段、状态、权限或数值。
2. 每条结论必须带 evidence_kind：
   - explicit：原文明确说明
   - inferred：从原文可以合理推断（必须写清推断依据，且不得变成新规则）
   - missing：原文未说明或存在歧义
3. 每条结论尽量带 sources（如 REQ-001），并引用原文短句 evidence。
4. 先识别/拆解需求条目，再做目标、功能、规则、角色、流程、数据、异常、缺口、风险。
5. 不要向用户连续提问。信息足够就完成分析；仅当缺口会改变分析结论时，最多给 1～3 个可跳过的关键问题。
6. 信息很模糊时，仍必须基于现有信息产出初步分析，不得拒绝工作。
7. 业务规则必须单独结构化，不要埋在长段落里。
8. 异常与边界是识别「需求是否定义了这些情况」，不是生成测试用例。
9. 数据字段若原文没说类型/必填/长度/默认值/唯一性，必须标 missing，禁止自行补具体规则。
10. 只输出一个 JSON 对象，不要 Markdown，不要解释。

JSON 结构：
{
  "requirements": [{"id":"REQ-001","text":"原文条目","source":"出现位置"}],
  "overview": {
    "name": {"text":"","evidence_kind":"explicit|inferred|missing","sources":["REQ-001"],"evidence":""},
    "goal": {"text":"","evidence_kind":"...","sources":[],"evidence":""},
    "background": {"text":"","evidence_kind":"...","sources":[],"evidence":""},
    "users": {"text":"","evidence_kind":"...","sources":[],"evidence":""},
    "in_scope": {"text":"","evidence_kind":"...","sources":[],"evidence":""},
    "out_of_scope": {"text":"","evidence_kind":"...","sources":[],"evidence":""}
  },
  "functions": [{
    "id":"FN-001","module":"模块","name":"子功能","parent":"",
    "description":"","inputs":[],"preconditions":[],"operations":[],
    "system_behavior":[],"outputs":[],"rules":["BR-001"],
    "sources":["REQ-001"],"evidence_kind":"explicit","evidence":""
  }],
  "rules": [{
    "id":"BR-001","category":"条件限制|数值限制|时间限制|状态限制|唯一性规则|数据关联规则|权限规则|审批规则|状态流转规则|其他",
    "statement":"","condition":"","sources":["REQ-001"],"evidence_kind":"explicit","evidence":""
  }],
  "roles": [{
    "id":"ROLE-001","name":"","can":[],"cannot":[],"data_scope":"","operations":[],
    "difference":"","sources":[],"evidence_kind":"missing","evidence":""
  }],
  "flows": [{
    "id":"FLOW-001","name":"","kind":"normal|exception|branch|state",
    "steps":["开始","用户操作","系统校验","成功/失败","状态变化","后续操作"],
    "states":[],"sources":[],"evidence_kind":"inferred","evidence":""
  }],
  "data": [{
    "id":"DATA-001","object":"","field":"","type":"","required":"","length":"",
    "default":"","unique":"","relations":"","lifecycle":"",
    "sources":[],"evidence_kind":"missing","evidence":""
  }],
  "exceptions": [{
    "id":"EX-001","scenario":"空值|非法输入|超长输入|最大/最小值|重复数据|无权限|状态不允许操作|网络异常|接口失败|数据不存在|并发操作|重复提交|超时|第三方服务异常|其他",
    "defined": false,
    "requirement":"需求如何定义，或写「需求未定义」",
    "sources":[],"evidence_kind":"missing"
  }],
  "gaps": [{"id":"GAP-001","problem":"","impact":"","need_confirm":"","sources":[]}],
  "ambiguities": [{"id":"AMB-001","original":"","questions":["只问会影响测试设计的问题"],"sources":[]}],
  "risks": [{"id":"RSK-001","level":"高|中|低","point":"","reason":"","scope":"","focus":"","sources":[]}],
  "test_focus": ["可测试关注点，不是用例"],
  "questions": [{"id":"Q1","question":"关键问题","skippable":true}]
}

exceptions 必须主动检查常见异常是否被需求定义，未定义就 defined=false。
questions 最多 3 条；需求完整时返回空数组。
"""


def empty_document() -> dict[str, Any]:
    return {
        "status": "empty",
        "source_text": "",
        "source_files": [],
        "completeness": 0,
        "counts": empty_counts(),
        "interaction": {
            "mode": "empty",
            "summary": "还没有需求分析结果。粘贴 PRD 或上传需求材料后开始分析。",
            "questions": [],
        },
        "requirements": [],
        "overview": {},
        "functions": [],
        "rules": [],
        "roles": [],
        "flows": [],
        "data": [],
        "exceptions": [],
        "gaps": [],
        "ambiguities": [],
        "risks": [],
        "test_focus": [],
        "user_notes": "",
        "answers": [],
        "design_input": "",
        "llm_used": False,
        "llm_error": None,
        "updated_at": None,
        "confirmed_at": None,
    }


def empty_counts() -> dict[str, int]:
    return {
        "requirements": 0,
        "functions": 0,
        "rules": 0,
        "open_questions": 0,
        "high_risks": 0,
    }


def _now() -> str:
    return datetime.now(timezone.utc).isoformat()


def _text_of(value: Any) -> str:
    if isinstance(value, dict):
        return str(value.get("text") or "").strip()
    return str(value or "").strip()


def _field_filled(value: Any) -> bool:
    text = _text_of(value)
    return bool(text) and text not in {"需求缺失", "未说明", "原文未说明"}


def parse_json_object(raw: str) -> dict[str, Any]:
    text = (raw or "").strip()
    if text.startswith("```"):
        text = re.sub(r"^```(?:json)?\s*", "", text, flags=re.I)
        text = re.sub(r"\s*```$", "", text)
    start = text.find("{")
    end = text.rfind("}")
    if start < 0 or end <= start:
        raise ValueError("模型没有返回 JSON")
    return json.loads(text[start : end + 1])


def extract_file_text(filename: str, payload: bytes) -> str:
    name = (filename or "file").lower()
    if name.endswith(".pdf"):
        try:
            import pdfplumber

            with pdfplumber.open(io.BytesIO(payload)) as pdf:
                return "\n".join((page.extract_text() or "") for page in pdf.pages)
        except Exception as exc:
            return f"[未能提取 PDF「{filename}」：{exc}]"
    if name.endswith(".docx"):
        try:
            from docx import Document

            doc = Document(io.BytesIO(payload))
            parts = [p.text for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if cells:
                        parts.append(" | ".join(cells))
            return "\n".join(parts)
        except Exception as exc:
            return f"[未能提取 Word「{filename}」：{exc}]"
    if name.endswith(".doc"):
        return f"[暂不支持旧版 .doc「{filename}」，请另存为 .docx、PDF 或文本]"
    try:
        return payload.decode("utf-8")
    except UnicodeDecodeError:
        return payload.decode("gbk", errors="replace")


def strip_agent_tags(text: str) -> str:
    raw = (text or "").replace("\r\n", "\n")
    lines = raw.split("\n")
    if lines and lines[0].startswith("@") and all(part.startswith("@") or not part.strip() for part in re.split(r"\s+", lines[0]) if part):
        return "\n".join(lines[1:]).strip()
    return raw.strip()


def split_requirements(text: str) -> list[dict[str, Any]]:
    raw = (text or "").strip()
    if not raw:
        return []
    items: list[dict[str, Any]] = []
    labeled = re.findall(
        r"(REQ[-_]?\d+)\s*[:：\.、]?\s*(.+?)(?=(?:\n\s*REQ[-_]?\d+)|$)",
        raw,
        flags=re.I | re.S,
    )
    if labeled:
        for rid, body in labeled:
            text_item = re.sub(r"\s+", " ", body).strip()
            if text_item:
                items.append({"id": rid.upper().replace("_", "-"), "text": text_item[:500], "source": "原文标注"})
        return items[:80]
    numbered = re.findall(r"(?m)^\s*(?:\d+[\.、\)］)]|[-*•])\s+(.+)$", raw)
    chunks = numbered if len(numbered) >= 2 else [p.strip() for p in re.split(r"\n\s*\n", raw) if p.strip()]
    if len(chunks) == 1 and len(raw) > 80:
        sentences = [s.strip() for s in re.split(r"(?<=[。！？；\n])", raw) if s.strip()]
        chunks = sentences if len(sentences) >= 2 else chunks
    for index, chunk in enumerate(chunks[:80], start=1):
        text_item = re.sub(r"\s+", " ", chunk).strip()
        if not text_item:
            continue
        items.append({"id": f"REQ-{index:03d}", "text": text_item[:500], "source": "原文拆解"})
    return items


def compute_counts(doc: dict[str, Any]) -> dict[str, int]:
    gaps = [item for item in (doc.get("gaps") or []) if not item.get("resolved")]
    ambiguities = doc.get("ambiguities") or []
    questions = (doc.get("interaction") or {}).get("questions") or doc.get("questions") or []
    open_questions = len(gaps) + len(ambiguities)
    if not open_questions:
        open_questions = len(questions)
    return {
        "requirements": len(doc.get("requirements") or []),
        "functions": len(doc.get("functions") or []),
        "rules": len(doc.get("rules") or []),
        "open_questions": open_questions,
        "high_risks": len([item for item in (doc.get("risks") or []) if (item.get("level") or "") == "高"]),
    }


def compute_completeness(doc: dict[str, Any]) -> int:
    overview = doc.get("overview") or {}
    filled = sum(
        1
        for key in ("name", "goal", "background", "users", "in_scope")
        if _field_filled(overview.get(key))
    )
    score = int(15 * filled / 5)
    if doc.get("requirements"):
        score += 10
    if doc.get("functions"):
        score += 15
    if doc.get("rules"):
        score += 10
    if doc.get("roles"):
        score += 10
    if doc.get("flows"):
        score += 10
    if any(_field_filled(item.get("field") or item.get("object")) for item in (doc.get("data") or [])):
        score += 10
    if doc.get("exceptions"):
        score += 8
    if "gaps" in doc:
        score += 6
    if doc.get("risks"):
        score += 6
    open_q = compute_counts(doc)["open_questions"]
    score -= min(20, open_q * 3)
    return max(0, min(100, score))


def interaction_of(doc: dict[str, Any]) -> dict[str, Any]:
    counts = compute_counts(doc)
    questions = []
    for item in (doc.get("questions") or [])[:3]:
        question = (item.get("question") if isinstance(item, dict) else str(item)).strip()
        if question:
            questions.append({
                "id": item.get("id") if isinstance(item, dict) else f"Q{len(questions) + 1}",
                "question": question,
                "skippable": True,
            })
    if not questions:
        for item in (doc.get("ambiguities") or [])[:3]:
            qs = item.get("questions") or []
            text = qs[0] if qs else item.get("original") or ""
            if text:
                questions.append({"id": item.get("id") or f"Q{len(questions) + 1}", "question": text, "skippable": True})
    if not questions:
        for item in (doc.get("gaps") or [])[:3]:
            text = item.get("need_confirm") or item.get("problem") or ""
            if text:
                questions.append({"id": item.get("id") or f"Q{len(questions) + 1}", "question": text, "skippable": True})
    questions = questions[:3]
    fn = counts["functions"]
    rules = counts["rules"]
    risks = len(doc.get("risks") or [])
    open_q = counts["open_questions"]
    if fn >= 3 and open_q == 0:
        mode = "complete"
        summary = f"已完成需求分析，共识别 {fn} 项功能、{rules} 条业务规则、{risks} 个风险项。"
    elif open_q <= 3:
        mode = "few_gaps"
        summary = f"已完成需求分析，共识别 {fn} 项功能、{rules} 条业务规则。另有 {open_q} 个关键问题，补充后可以提高准确度，也可以直接跳过。"
    else:
        mode = "vague"
        summary = f"已基于现有信息完成初步分析，共识别 {fn} 项功能、{rules} 条业务规则。当前需求存在 {open_q} 个关键缺口，补充后可以进一步提高分析准确度。"
    return {"mode": mode, "summary": summary, "questions": questions}


def to_design_input(doc: dict[str, Any]) -> str:
    if not doc or doc.get("status") in {None, "empty"}:
        return ""
    overview = doc.get("overview") or {}
    lines = [
        "【已确认的需求分析结果，供测试设计使用，不是测试用例】",
        f"需求名称：{_text_of(overview.get('name'))}",
        f"需求目标：{_text_of(overview.get('goal'))}",
        f"范围：{_text_of(overview.get('in_scope'))}",
        f"非范围：{_text_of(overview.get('out_of_scope'))}",
        "",
        "原始需求：",
    ]
    for item in doc.get("requirements") or []:
        lines.append(f"- {item.get('id')}: {item.get('text')}")
    lines.append("")
    lines.append("功能：")
    for item in doc.get("functions") or []:
        lines.append(
            f"- {item.get('id')} [{item.get('module')}/{item.get('name')}] "
            f"{item.get('description') or ''} 来源:{','.join(item.get('sources') or [])}"
        )
    lines.append("")
    lines.append("业务规则：")
    for item in doc.get("rules") or []:
        lines.append(
            f"- {item.get('id')} ({item.get('category')}) {item.get('statement')} "
            f"来源:{','.join(item.get('sources') or [])}"
        )
    lines.append("")
    lines.append("角色权限：")
    for item in doc.get("roles") or []:
        lines.append(f"- {item.get('name')} 可做:{'；'.join(item.get('can') or [])} 不可做:{'；'.join(item.get('cannot') or [])}")
    lines.append("")
    lines.append("待确认：")
    for item in doc.get("gaps") or []:
        if not item.get("resolved"):
            lines.append(f"- {item.get('id')}: {item.get('problem')} → {item.get('need_confirm')}")
    for item in doc.get("ambiguities") or []:
        lines.append(f"- {item.get('id')}: {item.get('original')} / {'；'.join(item.get('questions') or [])}")
    lines.append("")
    lines.append("风险：")
    for item in doc.get("risks") or []:
        lines.append(f"- [{item.get('level')}] {item.get('point')} 关注:{item.get('focus')}")
    lines.append("")
    lines.append("测试关注点：")
    for item in doc.get("test_focus") or []:
        lines.append(f"- {item}")
    return "\n".join(lines).strip()


def _normalize_evidence(item: dict[str, Any]) -> dict[str, Any]:
    kind = item.get("evidence_kind") or "missing"
    if kind not in EVIDENCE_KINDS:
        kind = "missing"
    item["evidence_kind"] = kind
    if not isinstance(item.get("sources"), list):
        item["sources"] = [str(item.get("sources"))] if item.get("sources") else []
    return item


def normalize_document(raw: dict[str, Any], source_text: str, source_files: list[dict[str, Any]], previous: Optional[dict[str, Any]] = None) -> dict[str, Any]:
    doc = empty_document()
    prev = previous or {}
    doc["source_text"] = (source_text or prev.get("source_text") or "")[:MAX_SOURCE_CHARS]
    doc["source_files"] = source_files or prev.get("source_files") or []
    doc["user_notes"] = (raw.get("user_notes") if raw.get("user_notes") is not None else prev.get("user_notes")) or ""
    doc["answers"] = raw.get("answers") or prev.get("answers") or []
    requirements = raw.get("requirements") or split_requirements(doc["source_text"])
    if not requirements and doc["source_text"].strip():
        requirements = split_requirements(doc["source_text"])
    doc["requirements"] = []
    for index, item in enumerate(requirements[:80], start=1):
        if not isinstance(item, dict):
            item = {"text": str(item)}
        rid = str(item.get("id") or f"REQ-{index:03d}").upper()
        if not re.match(r"REQ-\d+", rid):
            rid = f"REQ-{index:03d}"
        text = str(item.get("text") or "").strip()
        if not text:
            continue
        doc["requirements"].append({"id": rid, "text": text[:500], "source": item.get("source") or "原文"})
    overview = raw.get("overview") or {}
    normalized_overview = {}
    for key in ("name", "goal", "background", "users", "in_scope", "out_of_scope"):
        value = overview.get(key)
        if isinstance(value, dict):
            normalized_overview[key] = _normalize_evidence({
                "text": _text_of(value),
                "evidence_kind": value.get("evidence_kind") or "missing",
                "sources": value.get("sources") or [],
                "evidence": value.get("evidence") or "",
            })
        else:
            text = str(value or "").strip()
            normalized_overview[key] = {
                "text": text,
                "evidence_kind": "inferred" if text else "missing",
                "sources": [],
                "evidence": "",
            }
    doc["overview"] = normalized_overview
    for key in ("functions", "rules", "roles", "flows", "data", "exceptions", "gaps", "ambiguities", "risks"):
        items = raw.get(key) if isinstance(raw.get(key), list) else []
        cleaned = []
        for index, item in enumerate(items, start=1):
            if not isinstance(item, dict):
                continue
            row = dict(item)
            if key == "functions":
                row.setdefault("id", f"FN-{index:03d}")
                row.setdefault("module", "")
                row.setdefault("name", "")
                for field in ("inputs", "preconditions", "operations", "system_behavior", "outputs", "rules", "sources"):
                    if not isinstance(row.get(field), list):
                        row[field] = [row[field]] if row.get(field) else []
            elif key == "rules":
                row.setdefault("id", f"BR-{index:03d}")
                if not row.get("statement"):
                    continue
            elif key == "roles":
                row.setdefault("id", f"ROLE-{index:03d}")
                for field in ("can", "cannot", "operations", "sources"):
                    if not isinstance(row.get(field), list):
                        row[field] = [row[field]] if row.get(field) else []
            elif key == "flows":
                row.setdefault("id", f"FLOW-{index:03d}")
                if not isinstance(row.get("steps"), list):
                    row["steps"] = [row["steps"]] if row.get("steps") else []
                kind = row.get("kind") or "normal"
                row["kind"] = kind if kind in {"normal", "exception", "branch", "state"} else "normal"
            elif key == "data":
                row.setdefault("id", f"DATA-{index:03d}")
                for field in ("type", "required", "length", "default", "unique", "relations", "lifecycle"):
                    if not str(row.get(field) or "").strip():
                        row[field] = "需求缺失"
                        row["evidence_kind"] = "missing"
            elif key == "exceptions":
                row.setdefault("id", f"EX-{index:03d}")
                row["defined"] = bool(row.get("defined"))
                if not row.get("requirement"):
                    row["requirement"] = "需求未定义"
                    row["defined"] = False
                    row["evidence_kind"] = "missing"
            elif key == "gaps":
                row.setdefault("id", f"GAP-{index:03d}")
            elif key == "ambiguities":
                row.setdefault("id", f"AMB-{index:03d}")
                if not isinstance(row.get("questions"), list):
                    row["questions"] = [row["questions"]] if row.get("questions") else []
            elif key == "risks":
                row.setdefault("id", f"RSK-{index:03d}")
                level = row.get("level") or "中"
                row["level"] = level if level in RISK_LEVELS else "中"
            cleaned.append(_normalize_evidence(row))
        doc[key] = cleaned
    focus = raw.get("test_focus") or []
    doc["test_focus"] = [str(item).strip() for item in focus if str(item).strip()][:30]
    doc["questions"] = raw.get("questions") or []
    doc["interaction"] = interaction_of(doc)
    doc["counts"] = compute_counts(doc)
    doc["completeness"] = compute_completeness(doc)
    doc["design_input"] = to_design_input(doc)
    doc["llm_used"] = bool(raw.get("llm_used"))
    doc["llm_error"] = raw.get("llm_error")
    doc["status"] = raw.get("status") or prev.get("status") or "draft"
    if doc["status"] not in {"empty", "draft", "confirmed"}:
        doc["status"] = "draft"
    doc["updated_at"] = raw.get("updated_at") or _now()
    doc["confirmed_at"] = raw.get("confirmed_at") or prev.get("confirmed_at")
    return doc


def fallback_document(source_text: str, reason: str, source_files: Optional[list[dict[str, Any]]] = None) -> dict[str, Any]:
    requirements = split_requirements(source_text)
    overview_name = ""
    for line in (source_text or "").splitlines():
        line = line.strip()
        if line and not line.startswith("#") or line.startswith("#"):
            overview_name = line.lstrip("# ").strip()
            if overview_name:
                break
    raw = {
        "requirements": requirements,
        "overview": {
            "name": {"text": overview_name[:80], "evidence_kind": "inferred" if overview_name else "missing", "sources": [requirements[0]["id"]] if requirements else [], "evidence": overview_name[:80]},
            "goal": {"text": "", "evidence_kind": "missing", "sources": [], "evidence": ""},
            "background": {"text": "", "evidence_kind": "missing", "sources": [], "evidence": ""},
            "users": {"text": "", "evidence_kind": "missing", "sources": [], "evidence": ""},
            "in_scope": {"text": "", "evidence_kind": "missing", "sources": [], "evidence": ""},
            "out_of_scope": {"text": "", "evidence_kind": "missing", "sources": [], "evidence": ""},
        },
        "functions": [],
        "rules": [],
        "roles": [],
        "flows": [],
        "data": [],
        "exceptions": [
            {"scenario": name, "defined": False, "requirement": "需求未定义", "evidence_kind": "missing"}
            for name in ("空值", "非法输入", "超长输入", "重复数据", "无权限", "状态不允许操作", "接口失败", "超时")
        ],
        "gaps": [{
            "id": "GAP-001",
            "problem": "当前仅完成原文拆解，尚未得到完整结构化分析",
            "impact": "功能、规则、权限和状态可能不完整，影响后续测试设计",
            "need_confirm": "请确认需求材料是否完整，或在模型可用后重新分析",
            "sources": [item["id"] for item in requirements[:3]],
        }],
        "ambiguities": [],
        "risks": [{
            "id": "RSK-001",
            "level": "中",
            "point": "自动分析未完成",
            "reason": reason,
            "scope": "整份需求分析结果",
            "focus": "先核对原文条目，再补全规则与权限",
            "sources": [],
        }],
        "test_focus": [item["text"] for item in requirements[:8]],
        "questions": [{
            "id": "Q1",
            "question": "是否需要补充角色权限、状态流转和异常处理后再分析一次？",
            "skippable": True,
        }],
        "llm_used": False,
        "llm_error": reason,
        "status": "draft",
    }
    return normalize_document(raw, source_text, source_files or [])


def _run_async(coro):
    try:
        loop = asyncio.get_event_loop()
        if loop.is_running():
            with concurrent.futures.ThreadPoolExecutor(max_workers=1) as pool:
                return pool.submit(asyncio.run, coro).result()
        return loop.run_until_complete(coro)
    except RuntimeError:
        return asyncio.run(coro)


class RequirementAnalysisDocService:
    def __init__(self, memory: Optional[ProjectMemoryService] = None):
        self.memory = memory or ProjectMemoryService()

    def get(self, user_id: int, project_id: int) -> dict[str, Any]:
        row = self.memory.get_by_title(user_id, project_id, DOC_KIND, DOC_TITLE)
        if not row:
            return empty_document()
        try:
            raw = json.loads(row.get("content") or "")
            if not isinstance(raw, dict):
                return empty_document()
        except Exception:
            return empty_document()
        doc = normalize_document(raw, raw.get("source_text") or "", raw.get("source_files") or [], raw)
        extra = row.get("extra") or {}
        if extra.get("status"):
            doc["status"] = extra["status"]
        if extra.get("confirmed_at"):
            doc["confirmed_at"] = extra["confirmed_at"]
        doc["updated_at"] = row.get("updated_at") or doc.get("updated_at")
        doc["design_input"] = to_design_input(doc)
        return doc

    def analyze(
        self,
        user_id: int,
        project_id: int,
        source_text: str = "",
        source_files: Optional[list[dict[str, Any]]] = None,
        answers: Optional[list[dict[str, Any]]] = None,
        project_name: str = "",
    ) -> dict[str, Any]:
        previous = self.get(user_id, project_id)
        text = (source_text or "").strip()
        files = source_files or []
        if not text and previous.get("source_text"):
            text = previous["source_text"]
            files = files or previous.get("source_files") or []
        if not text:
            raise ValueError("请先提供需求材料（粘贴文本或上传文档）")
        if answers:
            extra = "\n".join(
                f"用户补充 {item.get('id') or ''}：{item.get('answer') or item.get('question') or ''}"
                for item in answers
                if item.get("answer")
            )
            if extra:
                text = f"{text}\n\n【用户补充】\n{extra}"
        text = text[:MAX_SOURCE_CHARS]
        try:
            raw = self._analyze_with_llm(text, project_name=project_name)
            raw["llm_used"] = True
            raw["llm_error"] = None
            raw["answers"] = answers or previous.get("answers") or []
            raw["status"] = "draft"
            doc = normalize_document(raw, text, files, previous)
        except Exception as exc:
            doc = fallback_document(text, str(exc), files)
            doc["answers"] = answers or previous.get("answers") or []
        self._save(user_id, project_id, doc)
        return doc

    def save_edits(self, user_id: int, project_id: int, patch: dict[str, Any]) -> dict[str, Any]:
        current = self.get(user_id, project_id)
        if current.get("status") == "empty" and not current.get("source_text"):
            raise ValueError("还没有可修改的需求分析结果")
        merged = dict(current)
        for key in ("overview", "functions", "rules", "roles", "flows", "data", "exceptions", "gaps", "ambiguities", "risks", "test_focus", "user_notes", "answers", "questions"):
            if key in patch:
                merged[key] = patch[key]
        if "status" in patch and patch["status"] in {"draft", "confirmed"}:
            merged["status"] = patch["status"]
        merged["status"] = "draft" if merged.get("status") == "empty" else merged.get("status") or "draft"
        merged["updated_at"] = _now()
        if merged["status"] != "confirmed":
            merged["confirmed_at"] = None
        doc = normalize_document(merged, merged.get("source_text") or "", merged.get("source_files") or [], current)
        self._save(user_id, project_id, doc)
        return doc

    def confirm(self, user_id: int, project_id: int) -> dict[str, Any]:
        current = self.get(user_id, project_id)
        if current.get("status") == "empty" or not current.get("requirements"):
            raise ValueError("请先完成需求分析再确认")
        current["status"] = "confirmed"
        current["confirmed_at"] = _now()
        current["updated_at"] = current["confirmed_at"]
        doc = normalize_document(current, current.get("source_text") or "", current.get("source_files") or [], current)
        self._save(user_id, project_id, doc)
        self.memory.remember(
            user_id, project_id,
            kind="test_design",
            title="需求分析（已确认）",
            content=doc["design_input"],
            workspace="design",
            extra={"from": "requirement_analysis", "completeness": doc["completeness"]},
        )
        self.memory.remember(
            user_id, project_id,
            kind="confirmed",
            title="需求分析已确认",
            content=doc["interaction"]["summary"],
            workspace=DOC_WORKSPACE,
            extra={"completeness": doc["completeness"], "counts": doc["counts"]},
        )
        return doc

    def analyze_from_question(self, user_id: int, project_id: int, question: str, project_name: str = "") -> dict[str, Any]:
        text = strip_agent_tags(question)
        reuse_hints = ("分析当前需求", "分析这份需求", "重新分析", "需求分析", "分析需求")
        if len(text) < 80 and any(hint in text for hint in reuse_hints):
            current = self.get(user_id, project_id)
            if not current.get("source_text"):
                raise ValueError("当前项目还没有需求材料。请先在需求分析页粘贴或上传 PRD。")
            text = current["source_text"]
            return self.analyze(user_id, project_id, text, current.get("source_files") or [], project_name=project_name)
        return self.analyze(user_id, project_id, text, project_name=project_name)

    def _save(self, user_id: int, project_id: int, doc: dict[str, Any]) -> None:
        payload = dict(doc)
        self.memory.remember(
            user_id, project_id,
            kind=DOC_KIND,
            title=DOC_TITLE,
            content=json.dumps(payload, ensure_ascii=False),
            workspace=DOC_WORKSPACE,
            extra={
                "status": doc.get("status"),
                "completeness": doc.get("completeness"),
                "counts": doc.get("counts"),
                "confirmed_at": doc.get("confirmed_at"),
            },
        )

    def _analyze_with_llm(self, source_text: str, project_name: str = "") -> dict[str, Any]:
        requirements = split_requirements(source_text)
        numbered = "\n".join(f"{item['id']}: {item['text']}" for item in requirements[:40])
        user_prompt = (
            f"项目名称：{project_name or '当前项目'}\n"
            "下面是需求原文。请先认同或修正 REQ 拆解，再输出完整 JSON。\n\n"
            f"【已拆解的原文条目】\n{numbered or '（未能稳定切分，请按原文自行编号 REQ-001 起）'}\n\n"
            f"【需求原文】\n{source_text}"
        )

        async def _call() -> str:
            gateway = get_gateway()
            return await gateway.chat(
                agent_name="requirement_analysis",
                system_prompt=SYSTEM_PROMPT,
                user_prompt=user_prompt,
                temperature=0.2,
                max_tokens=8192,
                timeout=120,
            )

        content = _run_async(_call())
        return parse_json_object(content)
