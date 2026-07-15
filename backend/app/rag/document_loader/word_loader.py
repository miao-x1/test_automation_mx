"""Word 文档解析器"""
import logging
from docx import Document as DocxDocument
from app.rag.document_loader.base import BaseDocumentLoader
from app.rag.models import LoadedDocument, DocumentType

logger = logging.getLogger(__name__)

class WordLoader(BaseDocumentLoader):
    def __init__(self):
        super().__init__()
        self.supported_types = [DocumentType.WORD]
    
    def can_handle(self, file_path: str, mime_type: str = "") -> bool:
        ext = file_path.lower().split('.')[-1]
        return ext in ('doc', 'docx') or 'word' in mime_type.lower()
    
    async def load(self, file_path: str, **kwargs) -> LoadedDocument:
        doc = DocxDocument(file_path)
        paragraphs = []
        tables = []
        for para in doc.paragraphs:
            if para.text.strip():
                style = para.style.name if para.style else "Normal"
                paragraphs.append(f"{'#' if 'Heading' in style else ''} {para.text}")
        for table in doc.tables:
            rows = []
            for row in table.rows:
                rows.append([cell.text for cell in row.cells])
            tables.append({"data": rows})
        content = "\n".join(paragraphs)
        return LoadedDocument(
            source_type=DocumentType.WORD,
            file_path=file_path,
            file_name=file_path.split('/')[-1].split('\\')[-1],
            content=content,
            tables=tables,
            metadata={"paragraph_count": len(paragraphs), "table_count": len(tables)},
        )
